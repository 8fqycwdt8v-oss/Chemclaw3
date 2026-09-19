"""What happens to a replica when a parse does not come back.

`agent/attachments.py` caps how many uploads may be parsed at once and sheds the rest, and that cap
was real. What it could not do was ever get a slot *back*: a slot stands for a running worker
thread and is released by that thread's completion callback, so a parse that never returns held its
slot for the life of the process. Driven before the fix, at the shipped cap of 2: both callers were
freed at their timeout, `in_flight` was still 2 five seconds later, and every later upload was shed.
The upload path of that replica was down permanently and nothing said so.

These tests hold the fix — `ingest/documents/isolate.py` — from both ends: that the work really
leaves this process, that a parser's own refusal still arrives as itself across that boundary, and
that a parse past its deadline costs a slot for the deadline rather than for the process. The last
one is the regression test for the wedge, and it is written as "the *next* upload is served"
because that is the property a chemist experiences; `in_flight` is the mechanism, not the promise.

The fourth is one layer down and was found by driving the third: `core/netguard.py` refused the
`AF_UNIX` socket the forkserver talks to itself over, while its own docstring said it exempted
local IPC. Nothing asserted it in either direction, and the C half of the same control
(`netguard_preload.c`) had it right — it reads `sa_family` and checks only the internet families.
"""

import asyncio
import io
import socket
import subprocess
import sys
import textwrap
import time
import zipfile
from pathlib import Path

import pytest

from chemclaw.agent import attachments
from chemclaw.agent.attachments import (
    AttachmentError,
    AttachmentUnavailable,
    parse_attachment_off_loop,
)
from chemclaw.core import netguard
from chemclaw.core.config import settings
from chemclaw.core.metrics import METRICS
from chemclaw.ingest.documents import isolate
from chemclaw.ingest.documents.isolate import (
    ParseWorkerLost,
    parse_context,
    parse_document_isolated,
)
from chemclaw.ingest.documents.parse import (
    DocumentParseError,
    ScannedDocumentError,
    parse_document,
)
from tests.egress_probe import egress_posture
from tests.test_document_formats import _blank_pdf_bytes  # type: ignore[attr-defined]

# A CSV big enough that parsing it is unmistakably longer than the deadline the wedge test sets,
# and small enough that building it costs nothing.
#
# **Wide rows rather than many narrow ones, and that is a memory shape rather than a preference.**
# The same 20 MB as `b"aaaa,bbbb,cccc,dddd\n" * 1_000_000`, which is what this was, and which now
# exceeds `document_parse_memory_bytes` and comes back as a refusal instead of a slow parse: a
# million rows is a million `str` objects in the rendered lines, and a `str` costs ~50 bytes of
# header before its characters. 100,000 rows of the same total length cost a tenth of that.
# Re-measured on this tree after the change: 0.56 s isolated against the 0.2 s deadline below, a
# factor of 2.8 where this comment used to claim ten. Ten is no longer available and that is the
# budget working: a CSV slow enough for it is a CSV whose rendered text does not fit one parse.
_SLOW_CSV = (b",".join([b"a" * 24] * 8) + b"\n") * 100_000


def test_a_parse_child_is_still_inside_the_no_egress_posture() -> None:
    """A parse now runs in a process this repository did not previously have, so say what it may do.

    The whole deployment posture is that nothing dials out. Moving untrusted bytes into a *new*
    process is exactly the move that could carry them outside a guard armed in the parent, and
    "the child inherits it" is an assumption rather than an observation — `forkserver` starts its
    server by fork **and exec**, so the child's guard is whatever that fresh interpreter armed, not
    a copy of the parent's memory.

    It is armed, and the mechanism is worth naming because it is not obvious: `_PRELOAD` imports
    `chemclaw.ingest.documents.parse`, which imports `chemclaw.core.config`, whose module body ends
    in `arm_egress_guard(settings)`. So the guard is armed in the forkserver before it forks
    anything, and every parse child inherits an armed one.

    **The probe is in `tests/egress_probe.py` because the chain is what is under test.** A target
    `forkserver` pickles by reference is imported in the child along with its whole module, so
    while this probe lived here the child imported this file — and with it `chemclaw.core.config`,
    which armed the guard on the spot. Driven: with `_PRELOAD` emptied the test still passed. Its
    own module imports `socket` and `sys`, and `tests/__init__.py` imports nothing, so the only
    way the guard can be armed in that child is the chain this docstring names.
    """
    context = parse_context()
    reader, writer = context.Pipe(duplex=False)
    child = context.Process(target=egress_posture, args=(writer,))
    child.start()
    writer.close()
    try:
        outcome, armed = reader.recv()
    finally:
        reader.close()
        child.join()

    assert armed is True, "the egress guard is not armed inside a parse child"
    assert outcome == "refused", (
        f"a parse child's outbound connect was {outcome!r} rather than refused by the egress "
        "guard, so untrusted bytes are parsed in a process outside the no-egress posture"
    )


#: What one fork-and-answer round trip costs, as the floor a derived deadline may not go under.
#: Not a `Settings` field, for the reason `_REAP_SECONDS` in `ingest/documents/isolate.py` is not
#: one: it is the latency of process creation, not a posture. Measured at 13-17 ms over five round
#: trips on a loaded box; 20 ms is that measurement rounded up, and the margin is applied once at
#: the assertion rather than twice — padding the constant *and* multiplying it is what made the
#: first version of this guard refuse a deadline that was already five times the round trip.
_FORK_ROUND_TRIP_SECONDS = 0.02


def _in_process_parse_seconds(raw: bytes) -> float:
    """What `raw` costs to parse here, so a deadline can be derived instead of transcribed.

    In-process on purpose: the number wanted is the *parse*, and going through
    `parse_document_isolated` would fold a fork round trip into it and then be compared against a
    deadline that has to contain one.
    """
    started = time.perf_counter()
    parse_document("slow.csv", raw, None)
    return time.perf_counter() - started


def _warm_the_forkserver() -> None:
    """Pay the forkserver's one-off start before a test measures anything.

    Measured: the first isolated parse in a process is ~0.93 s (the server is exec'd and preloads
    the parsers) and every one after it is ~0.04 s under pytest, ~0.01 s from a plain entry point.
    A test that timed a *first* parse would be timing that start-up, which is exactly the number
    `attachment_parse_reap_grace_seconds` exists to absorb and not the number under test.
    """
    parse_document_isolated("warm.csv", b"a,b\n1,2\n", None, 60.0)


def test_the_parse_does_not_run_in_this_process(monkeypatch: pytest.MonkeyPatch) -> None:
    """The work crosses a process boundary, proven by breaking this process's copy of it.

    `isolate.parse_document` is the name `_parse_into` calls. Replacing it here would stop any
    in-process parse dead; the child imports its own copy from the forkserver, which never saw this
    assignment, so a parse that still succeeds could only have happened somewhere else.

    Asserted this way rather than by comparing pids because a pid would have to be smuggled back
    through the result — a test seam in production code to prove a property the code already has.
    """
    _warm_the_forkserver()

    def refuse(*args: object, **kwargs: object) -> None:
        raise AssertionError("the parse ran in the calling process")

    monkeypatch.setattr(isolate, "parse_document", refuse)
    parsed = parse_document_isolated("runs.csv", b"id,yield\nR-1,88\n", None, 60.0)
    assert parsed.rows == 1
    assert "R-1" in parsed.text


def test_a_parsers_refusal_crosses_the_process_boundary_as_itself() -> None:
    """A scanned PDF is still a `ScannedDocumentError`, not a generic "the reader died".

    This is what the tagged pair on the pipe buys. An exception raised in the child and left to
    kill it would reach the parent as an exit code, and the share sync — which counts scans apart
    from unsupported formats so an operator can see how much of a corpus needs OCR — would have
    been counting process deaths.
    """
    _warm_the_forkserver()
    with pytest.raises(ScannedDocumentError) as excinfo:
        parse_document_isolated("scan.pdf", _blank_pdf_bytes(), None, 60.0)
    assert "OCR" in str(excinfo.value)


def test_a_parse_past_its_deadline_is_killed_and_counted() -> None:
    """The direct form of the fix: the child is killed, and the kill is visible from a scrape."""
    _warm_the_forkserver()
    before = METRICS.value("chemclaw_document_parse_kills_total")
    started = time.monotonic()
    with pytest.raises(ParseWorkerLost):
        parse_document_isolated("slow.csv", _SLOW_CSV, None, 0.2)
    elapsed = time.monotonic() - started
    # The caller is freed on the deadline rather than on the parse: the whole point is that the
    # thread ends. Generous on the upper side because sending 20 MB to the child is real work.
    assert elapsed < 2.0, elapsed
    assert METRICS.value("chemclaw_document_parse_kills_total") - before == 1


async def test_a_parse_past_its_deadline_frees_its_slot_for_the_next_upload(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The regression test for the wedge, stated as what a chemist experiences.

    One slot, one upload that will not finish inside its deadline, and then an ordinary small file.
    Before `ingest/documents/isolate.py` the slot was held until the parse finished — so the second
    upload waited `attachment_parse_queue_seconds` and came back as a retryable 503 about a file
    that parses in milliseconds. Driven against that arrangement this test fails with
    `AttachmentUnavailable`, which is the shape of the production failure: uploads refused by a pod
    that has capacity on paper.

    Three of the four budgets are set against the fixture's *measured* parse rather than against
    each other, so that the defect cannot hide behind any of them. The queue wait is shorter than
    the parse: if it were longer the second upload would simply outwait the wedge. The caller's
    backstop is shorter than the parse too, and the shipped 5 s is not — with that in place an
    in-process parse finishes *inside* the backstop and comes back as a success, so the mutation
    would be caught by the wrong assertion and this test would not be evidence about slots at all.

    **The budgets were transcribed against a ~2.3 s parse, and that parse is no longer 2.3 s.**
    `D-2026-09-19-a-ceiling-on-the-archive-is-not-a-ceiling-on-the-parse` rewrote `_parse_csv` to
    render row by row, which cut this fixture to **0.36 s** — leaving the 0.2 s deadline below it by
    1.8x, so whether the parse outran its deadline came down to how fast the runner was. It passed
    three local full runs and CI on the commit that caused it, then failed CI with
    `DID NOT RAISE`. Making the fixture slow again is not available: measured, 15x the cell count
    buys 0.40 s to 0.73 s, because the cost is in the bytes rather than in the cells, and the size
    that would buy 2.3 s is now refused by `document_parse_memory_bytes`.

    So the deadline is derived from the fixture instead of written down beside it: a quarter of what
    the parse actually costs *here*, which holds its ratio on a runner of any speed. The floor is
    the fork round trip — measured at 17 ms with this box loaded — because the *small* upload has to
    fit inside the same deadline, and if the parser ever gets fast enough to squeeze those together
    this fails saying so rather than going quietly marginal again.
    """
    _warm_the_forkserver()
    cost = _in_process_parse_seconds(_SLOW_CSV)
    deadline = cost / 4
    assert deadline >= 3 * _FORK_ROUND_TRIP_SECONDS, (
        f"_SLOW_CSV now parses in {cost:.3f}s, so a deadline it overruns four times over is "
        f"{deadline:.3f}s — inside the {_FORK_ROUND_TRIP_SECONDS:.3f}s a fork round trip costs, "
        "which would time the *small* upload out as well and make this test pass for the wrong "
        "reason. The fixture has to get slower, or this stops being about the deadline."
    )
    monkeypatch.setattr(settings, "attachment_max_concurrent_parses", 1)
    monkeypatch.setattr(settings, "attachment_parse_timeout_seconds", deadline)
    monkeypatch.setattr(settings, "attachment_parse_reap_grace_seconds", deadline * 2.5)
    # Shorter than the parse, for the reason above: a queue wait past it would let the second
    # upload simply outwait the wedge instead of being shed by it.
    monkeypatch.setattr(settings, "attachment_parse_queue_seconds", cost / 2)
    monkeypatch.setattr(settings, "attachment_max_bytes", len(_SLOW_CSV) + 1)

    with pytest.raises(AttachmentError):
        await parse_attachment_off_loop("slow.csv", _SLOW_CSV)
    # The release rides `Future.add_done_callback`, which asyncio dispatches with `call_soon`,
    # so it lands on the next turn of the loop rather than before this coroutine resumes.
    await asyncio.sleep(0)
    assert attachments._PARSE_SLOTS.in_flight == 0

    started = time.monotonic()
    parsed = await parse_attachment_off_loop("small.csv", b"id,yield\nR-1,88\n")
    assert parsed.rows == 1
    # Served, not queued: the slot was free when it arrived. Bounded well under the queue wait
    # so "it eventually got in" cannot pass for "it was never shed".
    assert time.monotonic() - started < settings.attachment_parse_queue_seconds


async def test_the_cap_still_sheds_when_the_slots_are_genuinely_busy(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The other direction, so the test above cannot pass by the cap having quietly stopped working.

    A cap that returns every slot immediately would satisfy "the next upload is served" and protect
    nothing. Here the one slot is held by a parse that is still inside its deadline, and the second
    upload must still be shed with the retryable refusal.
    """
    _warm_the_forkserver()
    monkeypatch.setattr(settings, "attachment_max_concurrent_parses", 1)
    monkeypatch.setattr(settings, "attachment_parse_timeout_seconds", 30.0)
    monkeypatch.setattr(settings, "attachment_parse_queue_seconds", 0.1)
    monkeypatch.setattr(settings, "attachment_max_bytes", len(_SLOW_CSV) + 1)

    holder = asyncio.create_task(parse_attachment_off_loop("slow.csv", _SLOW_CSV))
    # Let the holder take the slot before the second upload asks for one.
    while attachments._PARSE_SLOTS.in_flight == 0:
        await asyncio.sleep(0.01)
    with pytest.raises(AttachmentUnavailable):
        await parse_attachment_off_loop("second.csv", b"a,b\n1,2\n")
    await holder


def test_local_ipc_is_not_refused_as_egress(tmp_path: Path) -> None:
    """An `AF_UNIX` address names a path, and a path leaves this host by no route.

    `netguard._host_of` fell through to `host = address` for a non-tuple address, so a unix-socket
    connect arrived at `_check` as a hostname, failed `is_loopback_host` and was refused — while
    the same function's docstring said it returned None for exactly this case. Measured: every
    forkserver connect to `/tmp/pymp-*/listener-*` was logged as "outbound connection … is not on
    the allowlist" and raised `EgressForbidden`, which is why the isolated parse could not run at
    all until this was fixed.

    Asserted at both levels — the address reader, and a real `connect` through the armed guard —
    because the unit half alone would still pass if the guard were later re-plumbed to read the
    address somewhere else.
    """
    assert netguard._host_of("/tmp/pymp-abc/listener-def") is None
    assert netguard._host_of(b"\x00an-abstract-name") is None
    # An internet address is still read, in both spellings, so this is a narrowing and not a hole.
    assert netguard._host_of(("example.invalid", 443)) == "example.invalid"
    assert netguard._host_of((b"example.invalid", 443)) == "example.invalid"

    path = str(tmp_path / "probe.sock")
    with socket.socket(socket.AF_UNIX, socket.SOCK_STREAM) as server:
        server.bind(path)
        server.listen(1)
        with socket.socket(socket.AF_UNIX, socket.SOCK_STREAM) as client:
            client.connect(path)  # refused with EgressForbidden before the fix
            accepted, _ = server.accept()
            accepted.close()


def test_a_child_that_stalls_after_its_first_byte_still_frees_the_worker_thread() -> None:
    """The wedge this module exists to close, re-measured one stage later than it was fixed.

    **The regression.** `reader.poll(timeout)` was the only deadline, and `poll` returning True
    *consumes* it: `recv` then blocks until the whole pickled message arrives or the pipe reaches
    EOF, and `join()` had no timeout at all. Nothing killed the child on either path. Driven
    against the shipped `parse_document_isolated` before this commit, with a 1 s deadline and 15 s
    of patience:

    | child behaviour                        | worker thread |
    | writes a truncated message, then stops | **still alive** |
    | answers correctly, then does not exit  | **still alive** |

    A thread that never ends never releases its parse slot, so at the shipped cap of two such
    uploads the replica's upload path is down for the life of the process — the exact failure
    `ingest/documents/isolate.py` was written for, reappearing after the byte that satisfied the
    only bound. `agent/attachments.py`'s `wait_for` backstop cannot see it, because that wait is
    `shield`ed: it frees the caller while the thread it stands for runs on.

    **A subprocess, and `tests/parse_stalls.py` says why at length**: the only channel into a
    `forkserver` child is the server's preload list, and that server is a process-wide singleton
    another test in this file has already warmed. The probe drives the *shipped* parent code —
    same function, same worker-thread arrangement `agent/attachments.py` uses — and substitutes
    only what the child does.

    The third case is the grandchild: `Process.kill()` signals one pid, so a parse that shells out
    used to leave the grandchild burning CPU after the slot came back.
    """
    probe = subprocess.run(
        [sys.executable, "-c", "from tests.parse_stalls import main; main()"],
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    assert probe.returncode == 0, probe.stderr
    reported = {
        line.split("|")[0]: line.split("|")[1:] for line in probe.stdout.splitlines() if "|" in line
    }
    assert set(reported) == {"truncate.txt", "linger.txt", "grandchild.txt", "orphans"}, reported
    for case in ("truncate.txt", "linger.txt", "grandchild.txt"):
        assert reported[case][0] == "ended", (
            f"the worker thread for a child that {case} was still alive after 15 s, so its parse "
            f"slot is held for the life of the process: {reported[case]}"
        )
    # The one that answered correctly must still have been answered — a fix that turned every
    # lingering child into a refusal would pass the liveness assertion above and lose a parse.
    assert reported["linger.txt"][2] == "ParsedDocument", reported["linger.txt"]
    assert reported["orphans"][0] == "0", (
        f"{reported['orphans'][0]} grandchild process(es) outlived the kill that freed the slot"
    )


def _workbook_of_shared_strings(references: int, wide: bool) -> bytes:
    r"""A legal `.xlsx` whose text is many times its expanded size, optionally one code point wide.

    Written as raw OOXML rather than through `openpyxl` because the point is a property of the
    format that `openpyxl` will not produce: a *shared string* is stored once in the archive and
    referenced from as many cells as the sheet likes, so the text `_parse_xlsx` builds is the
    length of the string times the number of references while the archive grows by ~30 bytes each.
    Nothing here is crafted or dishonest — every `file_size` in the central directory is true,
    which is what `_refuse_a_bomb` reads, and the workbook opens in Excel.

    `wide` adds one more shared string holding a single astral code point, referenced from exactly
    one cell. Every other character is ASCII. CPython stores a `str` at the width of its widest code
    point, and `_parse_xlsx` ends in one document-wide `"\\n\\n".join(blocks)`, so that one cell
    quadruples the whole document.

    Args:
        references: How many cells point at the long shared string.
        wide: Whether one further cell holds a single astral code point.

    Returns:
        The workbook's bytes.
    """
    run = ("tetrahydrofuran-4-methoxybenzaldehyde-isolated-yield-" * 19)[:1000]
    strings = [run, "\U0001f9ea"] if wide else [run]
    per_row = 20
    rows = []
    for row in range(1, references // per_row + 1):
        cells = "".join(
            f'<c r="{chr(65 + column)}{row}" t="s"><v>0</v></c>' for column in range(per_row)
        )
        rows.append(f'<row r="{row}">{cells}</row>')
    if wide:
        rows.append(f'<row r="{len(rows) + 1}"><c r="A{len(rows) + 1}" t="s"><v>1</v></c></row>')
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    package = "http://schemas.openxmlformats.org/package/2006/relationships"
    document = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
        ),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{package}"><Relationship Id="rId1" '
            f'Type="{document}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
        ),
        "xl/workbook.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<workbook xmlns="{main}" xmlns:r="{document}"><sheets>'
            '<sheet name="runs" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{package}">'
            f'<Relationship Id="rId1" Type="{document}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{document}/sharedStrings" Target="sharedStrings.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{main}"><sheetData>{"".join(rows)}</sheetData></worksheet>'
        ),
        "xl/sharedStrings.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<sst xmlns="{main}" count="{references + int(wide)}" uniqueCount="{len(strings)}">'
            + "".join(f"<si><t>{one}</t></si>" for one in strings)
            + "</sst>"
        ),
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path, body in parts.items():
            archive.writestr(path, body)
    return buffer.getvalue()


def _declared_expansion(raw: bytes) -> int:
    """What `_refuse_a_bomb` reads out of the archive's central directory."""
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        return sum(item.file_size for item in archive.infolist())


def test_a_legal_upload_cannot_spend_more_than_the_parse_budget_declares() -> None:
    """Every ceiling on a document is a number in the archive; this is the one on the parse.

    The workbook here is legal by every bound upstream of the parse — well under
    `attachment_max_bytes` on the wire and under `document_max_expanded_bytes` expanded, with a
    central directory that tells the truth about both. It still asks for two hundred million
    characters, because a shared string is stored once and read from as many cells as the sheet
    has. Driven in a 1Gi memory cgroup holding the front door's measured 523 MiB idle pair, two
    concurrent parses of a workbook of this shape — 222,485 bytes on the wire, 5.9 MiB expanded,
    96.3 M characters — took the parent process with `SIGKILL`, exit 137: a pod OOMKill, every
    connected turn lost, from an upload nothing was entitled to refuse.

    What refuses it now is `document_parse_memory_bytes`, enforced by the kernel on the child that
    does the allocating, so nothing written in the archive can move it. Asserted as a refusal
    rather than as a memory reading because a memory reading of a process this one does not
    `waitpid` on is not available here: the parse child belongs to the forkserver, so
    `RUSAGE_CHILDREN` never sees it. The pod-level number is in
    `tests/test_deploy_chart.py::PARSE_MIB_PER_PARSE_BUDGET_MIB`.
    """
    raw = _workbook_of_shared_strings(200_000, wide=False)
    assert len(raw) < settings.attachment_max_bytes, "the fixture stopped being a legal upload"
    assert _declared_expansion(raw) < settings.document_max_expanded_bytes, (
        "the fixture stopped being legal by the expansion ceiling, which is the whole point of it"
    )
    with pytest.raises(DocumentParseError) as refusal:
        parse_document_isolated("runs.xlsx", raw, None, 120.0)
    assert "memory" in str(refusal.value), (
        "a document refused for its size must say so; the caller cannot act on 'could not be read'"
    )


def test_one_wide_code_point_does_not_multiply_what_a_parse_may_spend() -> None:
    r"""The same workbook, one astral character apart, and the budget is the same budget.

    `_parse_xlsx` ends in one document-wide `"\\n\\n".join(blocks)`, and CPython stores a `str` at
    the width of its widest code point — so a single emoji, superscript minus, `Å` or equilibrium
    arrow in one cell quadruples the whole extracted document. Measured on a real memory cgroup, a
    legal 1,089,493-byte upload at 63.4 MiB expanded charged the pod 236 MiB pure-ASCII and 500 MiB
    with one astral character in it, against a chart constant of 3.1 MiB per expanded MiB that
    predicted 197 for both.

    Both halves are asserted, and the first is why this is not simply "wide documents are refused":
    the ASCII twin must still parse, or the bound would have been bought by refusing everything.
    """
    narrow = _workbook_of_shared_strings(30_000, wide=False)
    wide = _workbook_of_shared_strings(30_000, wide=True)
    assert len(wide) - len(narrow) < 200, "the two fixtures must differ by one character, not more"

    parsed = parse_document_isolated("narrow.xlsx", narrow, None, 120.0)
    assert parsed.text.isascii() and len(parsed.text) > 30_000_000

    with pytest.raises(DocumentParseError) as refusal:
        parse_document_isolated("wide.xlsx", wide, None, 120.0)
    assert "memory" in str(refusal.value)


def _markup_heavy_docx(paragraphs: int, runs: int) -> bytes:
    """A legal Word report whose cost is its markup rather than its text.

    Every word its own styled run, which is what Word itself produces after tracked changes, mixed
    fonts, a spell-check language pass or a round-trip through another tool. `python-docx` builds an
    lxml DOM out of that markup, so the cost is in the elements and not in the characters — which is
    why no ceiling read out of the archive predicts it.
    """
    body = "".join(
        "<w:p><w:pPr><w:jc w:val='both'/></w:pPr>"
        + "".join(
            '<w:r><w:rPr><w:b/><w:color w:val="1F4E79"/><w:sz w:val="22"/></w:rPr>'
            f'<w:t xml:space="preserve">word{run} </w:t></w:r>'
            for run in range(runs)
        )
        + "</w:p>"
        for _ in range(paragraphs)
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}<w:sectPr/></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types"><Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
            'relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/'
            '2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/'
            '2006/relationships"/>',
        )
        archive.writestr("word/document.xml", document)
    return buffer.getvalue()


def test_a_document_stopped_by_the_budget_says_so_even_when_a_c_parser_reported_it() -> None:
    """The refusal a markup-heavy `.docx` earns, which used to say the document was malformed.

    **lxml reports its own allocation failure rather than letting CPython raise**, so the
    `except MemoryError` arm that names this ceiling never fired for the one format that most needs
    it. Measured on the shipped path before this: a 485,186-byte Word report — 2,000 paragraphs of
    200 styled runs, 2,979,999 characters of text, legal by every bound upstream — came back as
    `could not read report.docx: unknown error (<string>, line 0)`. That is not a missing reason, it
    is a wrong one: it tells a chemist their perfectly good report is broken at line 0, which is
    worse than the generic wording `too_large_to_read` exists to replace.

    `_at_ceiling` is what renames it, and both arms are asserted because either alone passes on the
    wrong implementation. A document that is *really* unreadable must keep its own message, or the
    fix is "call everything a memory problem" — driven, the two populations do not overlap: a
    parse stopped by the budget fails with 0.1 MiB of its allowance left, and a truncated archive
    fails with the whole 160 MiB unspent.
    """
    raw = _markup_heavy_docx(2_000, 200)
    assert len(raw) < settings.attachment_max_bytes, "the fixture stopped being a legal upload"
    assert _declared_expansion(raw) < settings.document_max_expanded_bytes, (
        "the fixture stopped being legal by the expansion ceiling, which is the point of it"
    )

    with pytest.raises(DocumentParseError) as refusal:
        parse_document_isolated("report.docx", raw, None, 120.0)
    assert "memory" in str(refusal.value), (
        f"a document the budget stopped was refused as {str(refusal.value)!r}, which reads as "
        "'your file is malformed' and names neither the ceiling nor the knob that moves it"
    )

    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as archive:
        archive.writestr("word/document.xml", "<w:document><not closed")
    with pytest.raises(DocumentParseError) as unreadable:
        parse_document_isolated("broken.docx", broken.getvalue(), None, 120.0)
    assert "memory" not in str(unreadable.value), (
        "a genuinely unreadable document was blamed on the memory budget, so the assertion above "
        "proves nothing — every refusal would pass it"
    )


def test_an_ambient_hard_limit_below_the_budget_is_a_smaller_budget_not_a_dead_parser() -> None:
    """`setrlimit` cannot raise a maximum, and that used to make every document unreadable.

    A process tree carrying any hard `RLIMIT_DATA` below `VmData + document_parse_memory_bytes` — a
    systemd `LimitDATA=`, a container security profile, an operator raising the knob above what the
    platform allows — made `_bound_allocations` raise `ValueError: not allowed to raise maximum
    limit`. In `_parse_into` that lands in the broad arm, so **every upload and every share document
    of every format** came back as "could not be read", with the cause only in a log line.

    Driven in a subprocess, because the limit has to be lowered before the call and a test process
    that lowers its own hard limit cannot put it back.
    """
    probe = textwrap.dedent(
        """
        import resource, sys
        from chemclaw.ingest.documents.isolate import _anonymous_bytes, _bound_allocations

        base = _anonymous_bytes()
        tight = base + 8 * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_DATA, (tight, tight))
        ceiling, hard = _bound_allocations(160 * 1024 * 1024)
        print("CLAMPED" if ceiling == tight and hard == tight else f"UNEXPECTED:{ceiling},{hard}")
        """
    )
    result = subprocess.run(
        [sys.executable, "-c", probe], capture_output=True, text=True, timeout=120, check=False
    )
    assert result.returncode == 0, (
        "a hard RLIMIT_DATA below the parse budget took the bound out with an exception, so every "
        f"document of every format is refused as unreadable: {result.stderr[-400:]}"
    )
    assert "CLAMPED" in result.stdout, (
        "the ambient ceiling did not become the budget; a lower platform limit is a smaller "
        f"budget, which is what this knob is for: {result.stdout!r}"
    )


def _markup_heavy_pptx(slides: int, runs: int) -> bytes:
    """A legal deck whose cost is its markup, the `.pptx` twin of `_markup_heavy_docx`.

    Every word its own styled run, which is what a deck becomes after a template change or a
    round-trip through another tool. `python-pptx` builds the same lxml DOM `python-docx` does, so
    the cost is in the elements rather than in the characters.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    deck = Presentation()
    blank = deck.slide_layouts[6]
    for _ in range(slides):
        slide = deck.slides.add_slide(blank)
        frame = slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(9), Inches(6)).text_frame
        paragraph = frame.paragraphs[0]
        for index in range(runs):
            run = paragraph.add_run()
            run.text = f"word{index} "
            run.font.bold = True
            run.font.size = Pt(11)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_a_deck_stopped_by_the_budget_earns_the_same_named_refusal_a_document_does() -> None:
    """`.pptx` goes through the same lxml layer, and nothing had driven it.

    `D-2026-09-19-a-refusal-that-blames-the-document-is-worse-than-one-that-says-nothing` fixed the
    `.docx` case — lxml reports its own allocation failure, so the arms that name this ceiling
    never fired and a legal report was refused as malformed at line 0 — and left the deck path
    unverified, with a `BACKLOG.md` row saying so rather than a guess.

    Driven: a 1,552,596-byte deck of 600 slides x 600 styled runs holds 2,821,690 characters,
    parses unbounded in 4.4 s, and is refused here in 0.5 s **with the memory refusal**, so
    `_at_ceiling` already covered it. This is what keeps that true rather than incidental.

    Both arms, as for `.docx`: a deck that is really unreadable must keep its own message, or every
    refusal would pass the first assertion.
    """
    raw = _markup_heavy_pptx(600, 600)
    assert len(raw) < settings.attachment_max_bytes, "the fixture stopped being a legal upload"

    with pytest.raises(DocumentParseError) as refusal:
        parse_document_isolated("deck.pptx", raw, None, 120.0)
    assert "memory" in str(refusal.value), (
        f"a deck the budget stopped was refused as {str(refusal.value)!r}, which names neither the "
        "ceiling nor the knob that moves it"
    )

    broken = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<p:presentation><not closed")
    with pytest.raises(DocumentParseError) as unreadable:
        parse_document_isolated("broken.pptx", broken.getvalue(), None, 120.0)
    assert "memory" not in str(unreadable.value), (
        "a genuinely unreadable deck was blamed on the memory budget, so the assertion above "
        "proves nothing"
    )


def test_a_refusal_is_not_bounded_by_the_ceiling_that_caused_it() -> None:
    """The reply is released before it is sent, because a `MemoryError` in a handler reaches no arm.

    Python does not route an exception raised inside one `except` clause to a later one, so a
    `MemoryError` while pickling the refusal onto the pipe escapes `_parse_into` entirely and the
    caller gets an EOF it can only report as "stopped without answering" — the one failure in that
    module whose cause is knowable, arriving nameless. The refusal is ~250 characters, so it is
    unlikely rather than impossible, and "unlikely" is an argument rather than a measurement.

    `_release_allocations` removes the question instead of estimating it: by the time a handler
    runs the parse is over, the budget's job is done, and the reply gets the room the parse was
    denied. Asserted on the mechanism — the ceiling is gone once the arm has run — because
    provoking a real allocation failure inside a handler is not something a test can stage
    honestly.
    """
    probe = textwrap.dedent(
        """
        import resource
        from chemclaw.ingest.documents.isolate import _bound_allocations, _release_allocations

        before = resource.getrlimit(resource.RLIMIT_DATA)
        ceiling, hard = _bound_allocations(64 * 1024 * 1024)
        bounded = resource.getrlimit(resource.RLIMIT_DATA)[0]
        _release_allocations(hard)
        after = resource.getrlimit(resource.RLIMIT_DATA)[0]
        print("BOUNDED" if bounded == ceiling else f"NOT-BOUNDED:{bounded}")
        print("RELEASED" if after == hard and after != bounded else f"STILL-BOUND:{after}")
        same = resource.getrlimit(resource.RLIMIT_DATA)[1] == before[1]
        print("UNCHANGED-HARD" if same else "HARD-MOVED")
        """
    )
    result = subprocess.run(
        [sys.executable, "-c", probe], capture_output=True, text=True, timeout=120, check=False
    )
    assert result.returncode == 0, f"the probe did not run: {result.stderr[-400:]}"
    assert "BOUNDED" in result.stdout, f"the parse was never bounded: {result.stdout!r}"
    assert "RELEASED" in result.stdout, (
        "the ceiling still stood after the failure arm released it, so a refusal is pickled under "
        f"the budget that stopped the parse: {result.stdout!r}"
    )
    assert "UNCHANGED-HARD" in result.stdout, (
        "releasing moved the hard limit, which is irreversible for the process and is not what "
        f"this is for: {result.stdout!r}"
    )
